"""Extract checklist master data + self-assessment answers from a Marubeni ITGR workbook.

Replaces the one-off scripts/extract-data.mjs (which scraped the legacy index.html).
Reads the official .xlsx directly so future FY updates are a re-run, not a re-scrape.

  python scripts/extract-xlsx.py "<path to workbook.xlsx>"

Writes:
  data/items.json       — the 96 requirements (master data for checklist_items)
  data/assessment.json  — the self-assessment answers recorded in the workbook

The short category labels (catShort / catShortTh) are NOT in the workbook — they are
hand-authored bilingual labels used across the dashboard UI, so they are carried
forward from the existing data/items.json rather than regenerated.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent

# Column layout, verified identical across FY2025 and FY2026 workbooks.
COL = {
    "no": 1, "category": 2, "name": 3, "content": 4, "standard": 5,
    "evidence": 6, "article": 7, "issue": 8, "risk": 13, "priority": 14, "qtype": 15,
}
COL_CHECK = 19    # S — self-assessment answer: 〇 / × / −
COL_REMARKS = 20  # T — free-text justification
FIRST_DATA_ROW = 10


def norm(v):
    if v is None:
        return ""
    s = str(v).replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"[ \t　]+", " ", s)
    return "\n".join(line.strip() for line in s.split("\n")).strip()


def pick_sheet(wb):
    """Sheet name drifts between releases ('ITGR Checklist' vs 'Current_ITGR Checklist_EN ')."""
    for name in wb.sheetnames:
        if "checklist" in name.lower():
            return wb[name]
    raise SystemExit(f"No checklist sheet found. Sheets: {wb.sheetnames}")


def main():
    # Workbook titles carry CJK; the Windows console default codepage cannot encode them.
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python scripts/extract-xlsx.py <workbook.xlsx>")
    path = sys.argv[1]

    # Short labels are UI copy, not source data — carry them over.
    existing = json.loads((ROOT / "data/items.json").read_text(encoding="utf-8"))
    labels = {i["catNo"]: (i["catShort"], i["catShortTh"]) for i in existing}

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = pick_sheet(wb)

    items, assessment = [], []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        raw_no = ws.cell(r, COL["no"]).value
        if raw_no is None or not str(raw_no).strip().isdigit():
            continue
        no = int(raw_no)
        row = {k: norm(ws.cell(r, c).value) for k, c in COL.items()}
        cat_no = int(row["category"].split(".")[0])
        short, short_th = labels.get(cat_no, ("", ""))
        items.append({
            "no": no,
            "category": row["category"],
            "name": row["name"],
            "content": row["content"],
            "standard": row["standard"],
            "evidence": row["evidence"],
            "article": row["article"],
            "issue": row["issue"],
            "risk": row["risk"],
            "priority": row["priority"],
            "qtype": row["qtype"],
            "catNo": cat_no,
            "catShort": short,
            "catShortTh": short_th,
        })
        answer = norm(ws.cell(r, COL_CHECK).value)
        remarks = norm(ws.cell(r, COL_REMARKS).value)
        if answer or remarks:
            assessment.append({"no": no, "answer": answer, "remarks": remarks})

    (ROOT / "data/items.json").write_text(
        json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
    (ROOT / "data/assessment.json").write_text(
        json.dumps(assessment, ensure_ascii=False, indent=2), encoding="utf-8")

    title = norm(ws.cell(1, 1).value).replace("\n", " ")
    tally = {}
    for a in assessment:
        tally[a["answer"] or "(blank)"] = tally.get(a["answer"] or "(blank)", 0) + 1
    print(f"source : {title}")
    print(f"items  : {len(items)} -> data/items.json")
    print(f"answers: {len(assessment)} rows -> data/assessment.json  {tally}")


if __name__ == "__main__":
    main()
